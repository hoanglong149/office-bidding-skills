#!/usr/bin/env python3
"""
demo_project.py — Demo: lấy bộ biểu mẫu TT79 có sẵn trong `templates/EPC-10A/`,
chép ra thư mục làm việc và điền vài dòng dữ liệu mẫu vào các bảng Excel.

    python3 examples/demo_project.py [thư_mục_output]

Không cần dữ liệu thật — dùng để xem bộ biểu mẫu trông thế nào và học cách điền.
"""

from __future__ import annotations
import os
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TEMPLATES = os.path.join(ROOT, "templates", "EPC-10A")

# --------------------------------------------------------------- dữ liệu mẫu
# (cột theo đúng tiêu đề của Mẫu 06A / 06D / 10B trong bộ template)
NHAN_SU = [
    ("Nguyễn Văn A", "Chỉ huy trưởng công trường"),
    ("Trần Thị B", "Kỹ sư trưởng công nghệ"),
    ("Lê Văn C", "Kỹ sư trưởng điện – C&I"),
    ("Phạm Thị D", "Kỹ sư an toàn (HSE)"),
]

THIET_BI = [
    ("Cần cẩu bánh xích", "Liebherr", "LR 1600/2", "600 t", 2019),
    ("Máy hàn inverter", "Lincoln Electric", "Vantage 500", "500 A", 2022),
    ("Máy nén khí", "Atlas Copco", "XA186", "10 m³/min", 2021),
]

HANG_HOA = [
    ("Bồn chứa khí nguyên liệu", "ABC-2026-V01", "ABC", 2026, "Việt Nam", "Công ty Cổ phần Cơ khí ABC", 2, "Bồn"),
    ("Bơm ly tâm", "EBR-CW-320", "Ebara", 2025, "Việt Nam", "Ebara Việt Nam", 4, "Cái"),
    ("Van điều khiển", "SAM-DN100", "Samson", 2025, "Đức", "Samson AG", 12, "Bộ"),
]

ROW_CLEAR = 20


def _write_rows(path: str, start_row: int, rows) -> None:
    """Ghi dữ liệu vào sheet 1, bắt đầu từ `start_row`; xoá các dòng trống phía dưới."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    for r in range(start_row, start_row + ROW_CLEAR):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).value = None
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            ws.cell(start_row + i, j).value = v
    wb.save(path)


def fill_nhan_su(path: str) -> None:
    # Mẫu 06A: STT | Họ và Tên | Vị trí công việc — dữ liệu từ dòng 2
    _write_rows(path, 2, [(i, ten, vi_tri) for i, (ten, vi_tri) in enumerate(NHAN_SU, 1)])


def fill_thiet_bi(path: str) -> None:
    # Mẫu 06D: STT|Loại|NSX|Model|Công suất|Năm SX|Tính năng|Xuất xứ|Địa điểm|Tình hình|Nguồn — từ dòng 3
    _write_rows(path, 3, [
        (i, loai, nsx, model, cs, nam, "", "—", "Công trường", "Sẵn sàng", "Sở hữu của nhà thầu")
        for i, (loai, nsx, model, cs, nam) in enumerate(THIET_BI, 1)
    ])


def fill_hang_hoa(path: str) -> None:
    # Mẫu 10B: STT|Danh mục|Ký mã hiệu|Nhãn hiệu|Năm SX|Xuất xứ|Hãng SX|Cấu hình|ĐVT|Khối lượng|Mã HS|Đơn giá|Thành tiền — từ dòng 3
    _write_rows(path, 3, [
        (i, ten, ma, nhan, nam, xx, hang, "", dvt, sl, "", "", "")
        for i, (ten, ma, nhan, nam, xx, hang, sl, dvt) in enumerate(HANG_HOA, 1)
    ])


TARGETS = {
    "TT79-M06A.xlsx": fill_nhan_su,
    "TT79-M06D.xlsx": fill_thiet_bi,
    "TT79-M10B.xlsx": fill_hang_hoa,
}


def main() -> None:
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "demo_output")
    os.makedirs(out, exist_ok=True)

    n = 0
    for f in sorted(os.listdir(TEMPLATES)):
        if f.endswith((".docx", ".xlsx", ".md")):
            shutil.copy2(os.path.join(TEMPLATES, f), os.path.join(out, f))
            n += 1
    print(f"1/2  Đã chép {n} biểu mẫu TT79 → {out}")

    filled = 0
    for name, fn in TARGETS.items():
        p = os.path.join(out, name)
        if os.path.exists(p):
            fn(p)
            filled += 1
    print(f"2/2  Đã điền dữ liệu mẫu vào {filled} bảng Excel (06A nhân sự, 06D thiết bị, 10B hàng hóa)")
    print(f"\nMở thư mục để xem: {out}")


if __name__ == "__main__":
    main()

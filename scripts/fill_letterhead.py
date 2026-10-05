#!/usr/bin/env python3
"""
fill_letterhead.py — Điền khối tiêu đề (letterhead) + dữ liệu dự án vào bộ biểu mẫu HSDT.

Thay các placeholder trong .docx (paragraph + table + header/footer) và .xlsx (cell):

  [TÊN NHÀ THẦU / LIÊN DANH NHÀ THẦU]   -> --bidder
  [Địa chỉ · Điện thoại · Email]        -> --bidder-info
  [TÊN CHỦ ĐẦU TƯ / BÊN MỜI THẦU]       -> --employer
  [Địa chỉ · Điện thoại · Fax]          -> --employer-info
  [Địa danh]                            -> --place
  ……/……-HSDT                            -> --doc-no
  ngày …… tháng …… năm 20……             -> --date

Dùng:
  python3 fill_letterhead.py <thư_mục> --bidder "..." --employer "..." --place "Hà Nội" \
        --bidder-info "..." --employer-info "..." --doc-no "01/2026-HSDT" --date "ngày 05 tháng 10 năm 2026"
"""

from __future__ import annotations
import argparse
import glob
import os
import re

from docx import Document


REPL = {}


def _sub(text: str) -> str:
    if not text:
        return text
    for k, v in REPL.items():
        text = text.replace(k, v)
    return text


def fill_docx(path: str) -> int:
    doc = Document(path)
    n = 0

    def walk_paras(paras):
        nonlocal n
        for p in paras:
            if not p.runs:
                continue
            full = "".join(r.text for r in p.runs)
            new = _sub(full)
            if new != full:
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
                n += 1

    walk_paras(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                walk_paras(c.paragraphs)
                for nt in c.tables:
                    for r2 in nt.rows:
                        for c2 in r2.cells:
                            walk_paras(c2.paragraphs)
    for s in doc.sections:
        for part in (s.header, s.footer):
            walk_paras(part.paragraphs)
    doc.save(path)
    return n


def fill_xlsx(path: str) -> int:
    import openpyxl
    wb = openpyxl.load_workbook(path)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    new = _sub(c.value)
                    if new != c.value:
                        c.value = new
                        n += 1
    wb.save(path)
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder")
    ap.add_argument("--bidder", default="")
    ap.add_argument("--bidder-info", default="")
    ap.add_argument("--employer", default="")
    ap.add_argument("--employer-info", default="")
    ap.add_argument("--place", default="")
    ap.add_argument("--date", default="")
    ap.add_argument("--doc-no", default="")
    a = ap.parse_args()

    if a.bidder:
        REPL["[TÊN NHÀ THẦU / LIÊN DANH NHÀ THẦU]"] = a.bidder
    if a.bidder_info:
        REPL["[Địa chỉ · Điện thoại · Email]"] = a.bidder_info
    if a.employer:
        REPL["[TÊN CHỦ ĐẦU TƯ / BÊN MỜI THẦU]"] = a.employer
    if a.employer_info:
        REPL["[Địa chỉ · Điện thoại · Fax]"] = a.employer_info
    if a.place:
        REPL["[Địa danh]"] = a.place
    if a.doc_no:
        REPL["……/……-HSDT"] = a.doc_no
    if a.date:
        REPL["ngày …… tháng …… năm 20……"] = a.date

    total = 0
    for f in sorted(glob.glob(os.path.join(a.folder, "*.docx"))):
        k = fill_docx(f)
        print(f"  docx  {os.path.basename(f):52} {k} thay thế")
        total += k
    for f in sorted(glob.glob(os.path.join(a.folder, "*.xlsx"))):
        k = fill_xlsx(f)
        print(f"  xlsx  {os.path.basename(f):52} {k} thay thế")
        total += k
    print(f"Tổng: {total} thay thế")


if __name__ == "__main__":
    main()
